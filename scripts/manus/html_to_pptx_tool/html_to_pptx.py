#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
HTML to PPTX Converter
将HTML报告转换为可编辑的PPTX演示文稿
保留字体、颜色、布局、表格、图片等效果
"""

import os
import re
import sys
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor as RgbColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from copy import deepcopy
import colorsys


class HTMLToPPTXConverter:
    """HTML转PPTX转换器"""
    
    def __init__(self, html_path, output_path=None):
        self.html_path = html_path
        self.html_dir = os.path.dirname(os.path.abspath(html_path))
        self.output_path = output_path or html_path.replace('.html', '.pptx')
        
        # 幻灯片尺寸（宽屏16:9）
        self.slide_width = Inches(13.333)
        self.slide_height = Inches(7.5)
        
        # 内容区域边距
        self.margin_left = Inches(0.5)
        self.margin_right = Inches(0.5)
        self.margin_top = Inches(0.5)
        self.margin_bottom = Inches(0.5)
        
        # 内容宽度
        self.content_width = self.slide_width - self.margin_left - self.margin_right
        
        # 颜色映射
        self.colors = {
            'primary': RgbColor(52, 152, 219),      # #3498db
            'secondary': RgbColor(44, 62, 80),       # #2c3e50
            'success': RgbColor(39, 174, 96),        # #27ae60
            'warning': RgbColor(243, 156, 18),       # #f39c12
            'danger': RgbColor(231, 76, 60),         # #e74c3c
            'info': RgbColor(23, 162, 184),          # #17a2b8
            'light': RgbColor(248, 249, 250),        # #f8f9fa
            'dark': RgbColor(52, 73, 94),            # #34495e
            'white': RgbColor(255, 255, 255),
            'black': RgbColor(51, 51, 51),           # #333
            'gradient_start': RgbColor(102, 126, 234),  # #667eea
            'gradient_end': RgbColor(118, 75, 162),     # #764ba2
        }
        
        # 字体设置
        self.fonts = {
            'title': 'Microsoft YaHei',
            'body': 'Microsoft YaHei',
            'code': 'Consolas',
        }
        
        # 创建演示文稿
        self.prs = Presentation()
        self.prs.slide_width = self.slide_width
        self.prs.slide_height = self.slide_height
        
        # 解析HTML
        with open(html_path, 'r', encoding='utf-8') as f:
            self.soup = BeautifulSoup(f.read(), 'html.parser')
    
    def parse_color(self, color_str):
        """解析CSS颜色值"""
        if not color_str:
            return None
        
        color_str = color_str.strip().lower()
        
        # 处理hex颜色
        if color_str.startswith('#'):
            hex_color = color_str[1:]
            if len(hex_color) == 3:
                hex_color = ''.join([c*2 for c in hex_color])
            if len(hex_color) == 6:
                r = int(hex_color[0:2], 16)
                g = int(hex_color[2:4], 16)
                b = int(hex_color[4:6], 16)
                return RgbColor(r, g, b)
        
        # 处理rgb颜色
        rgb_match = re.match(r'rgb\s*\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)', color_str)
        if rgb_match:
            r, g, b = map(int, rgb_match.groups())
            return RgbColor(r, g, b)
        
        # 处理rgba颜色
        rgba_match = re.match(r'rgba\s*\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*,\s*[\d.]+\s*\)', color_str)
        if rgba_match:
            r, g, b = map(int, rgba_match.groups())
            return RgbColor(r, g, b)
        
        # 常用颜色名称
        color_names = {
            'white': RgbColor(255, 255, 255),
            'black': RgbColor(0, 0, 0),
            'red': RgbColor(255, 0, 0),
            'green': RgbColor(0, 128, 0),
            'blue': RgbColor(0, 0, 255),
        }
        
        return color_names.get(color_str)
    
    def add_blank_slide(self):
        """添加空白幻灯片"""
        blank_layout = self.prs.slide_layouts[6]  # 空白布局
        return self.prs.slides.add_slide(blank_layout)
    
    def add_shape_with_gradient(self, slide, left, top, width, height, start_color, end_color):
        """添加渐变背景形状"""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, top, width, height
        )
        shape.line.fill.background()
        
        # 设置渐变填充
        fill = shape.fill
        fill.gradient()
        fill.gradient_angle = 135
        fill.gradient_stops[0].color.rgb = start_color
        fill.gradient_stops[1].color.rgb = end_color
        
        return shape
    
    def add_text_box(self, slide, left, top, width, height, text, 
                     font_size=12, font_color=None, font_bold=False,
                     font_name=None, align=PP_ALIGN.LEFT, 
                     vertical_anchor=MSO_ANCHOR.TOP):
        """添加文本框"""
        textbox = slide.shapes.add_textbox(left, top, width, height)
        tf = textbox.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        
        p = tf.paragraphs[0]
        p.alignment = align
        
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = font_bold
        run.font.name = font_name or self.fonts['body']
        
        if font_color:
            run.font.color.rgb = font_color
        
        tf.paragraphs[0].space_before = Pt(0)
        tf.paragraphs[0].space_after = Pt(0)
        
        return textbox
    
    def add_rectangle(self, slide, left, top, width, height, fill_color=None, 
                      border_color=None, border_width=0):
        """添加矩形"""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, top, width, height
        )
        
        if fill_color:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill_color
        else:
            shape.fill.background()
        
        if border_color and border_width > 0:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(border_width)
        else:
            shape.line.fill.background()
        
        return shape
    
    def create_title_slide(self, title, subtitle=None):
        """创建标题页"""
        slide = self.add_blank_slide()
        
        # 添加渐变背景
        bg_shape = self.add_shape_with_gradient(
            slide, Inches(0), Inches(0),
            self.slide_width, self.slide_height,
            self.colors['gradient_start'], self.colors['gradient_end']
        )
        
        # 添加标题
        title_box = self.add_text_box(
            slide, self.margin_left, Inches(2.5),
            self.content_width, Inches(1.5),
            title, font_size=44, font_color=self.colors['white'],
            font_bold=True, align=PP_ALIGN.CENTER
        )
        
        # 添加副标题
        if subtitle:
            subtitle_box = self.add_text_box(
                slide, self.margin_left, Inches(4),
                self.content_width, Inches(1),
                subtitle, font_size=20, font_color=self.colors['white'],
                align=PP_ALIGN.CENTER
            )
        
        return slide
    
    def create_section_title_slide(self, title):
        """创建章节标题页"""
        slide = self.add_blank_slide()
        
        # 左侧蓝色条
        left_bar = self.add_rectangle(
            slide, Inches(0), Inches(0),
            Inches(0.3), self.slide_height,
            fill_color=self.colors['primary']
        )
        
        # 标题
        title_box = self.add_text_box(
            slide, Inches(0.8), Inches(3),
            self.content_width, Inches(1.5),
            title, font_size=36, font_color=self.colors['secondary'],
            font_bold=True, align=PP_ALIGN.LEFT
        )
        
        return slide
    
    def create_content_slide(self, title, content_elements):
        """创建内容页"""
        slide = self.add_blank_slide()
        
        # 顶部蓝色条
        top_bar = self.add_rectangle(
            slide, Inches(0), Inches(0),
            self.slide_width, Inches(0.1),
            fill_color=self.colors['primary']
        )
        
        # 标题
        title_box = self.add_text_box(
            slide, self.margin_left, Inches(0.3),
            self.content_width, Inches(0.8),
            title, font_size=28, font_color=self.colors['secondary'],
            font_bold=True, align=PP_ALIGN.LEFT
        )
        
        # 标题下划线
        underline = self.add_rectangle(
            slide, self.margin_left, Inches(0.95),
            Inches(2), Pt(3),
            fill_color=self.colors['primary']
        )
        
        return slide
    
    def add_table_to_slide(self, slide, table_data, left, top, width, col_widths=None):
        """添加表格到幻灯片"""
        if not table_data or len(table_data) == 0:
            return None
        
        rows = len(table_data)
        cols = len(table_data[0]) if table_data else 0
        
        if cols == 0:
            return None
        
        # 计算行高
        row_height = Inches(0.4)
        table_height = row_height * rows
        
        # 创建表格
        table = slide.shapes.add_table(rows, cols, left, top, width, table_height).table
        
        # 设置列宽
        if col_widths:
            for i, w in enumerate(col_widths):
                if i < cols:
                    table.columns[i].width = w
        else:
            col_width = width // cols
            for i in range(cols):
                table.columns[i].width = col_width
        
        # 填充数据
        for row_idx, row_data in enumerate(table_data):
            for col_idx, cell_data in enumerate(row_data):
                if col_idx >= cols:
                    continue
                    
                cell = table.cell(row_idx, col_idx)
                cell.text = str(cell_data.get('text', ''))
                
                # 设置单元格样式
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(11)
                para.font.name = self.fonts['body']
                
                # 表头样式
                if row_idx == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self.colors['primary']
                    para.font.color.rgb = self.colors['white']
                    para.font.bold = True
                else:
                    # 交替行颜色
                    if row_idx % 2 == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = self.colors['light']
                    else:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = self.colors['white']
                    para.font.color.rgb = self.colors['black']
                
                # 处理badge样式
                if cell_data.get('badge'):
                    badge_type = cell_data.get('badge_type', 'medium')
                    if badge_type == 'high':
                        para.font.color.rgb = self.colors['success']
                    elif badge_type == 'low':
                        para.font.color.rgb = self.colors['danger']
                    else:
                        para.font.color.rgb = self.colors['warning']
                    para.font.bold = True
        
        return table
    
    def add_image_to_slide(self, slide, image_path, left, top, width=None, height=None):
        """添加图片到幻灯片"""
        if not os.path.exists(image_path):
            # 尝试相对路径
            rel_path = os.path.join(self.html_dir, image_path)
            if os.path.exists(rel_path):
                image_path = rel_path
            else:
                print(f"Warning: Image not found: {image_path}")
                return None
        
        try:
            if width and height:
                pic = slide.shapes.add_picture(image_path, left, top, width, height)
            elif width:
                pic = slide.shapes.add_picture(image_path, left, top, width=width)
            elif height:
                pic = slide.shapes.add_picture(image_path, left, top, height=height)
            else:
                pic = slide.shapes.add_picture(image_path, left, top)
            return pic
        except Exception as e:
            print(f"Error adding image {image_path}: {e}")
            return None
    
    def add_stat_card(self, slide, left, top, width, height, number, label):
        """添加统计卡片"""
        # 卡片背景
        card = self.add_rectangle(
            slide, left, top, width, height,
            fill_color=self.colors['white'],
            border_color=RgbColor(225, 232, 237),
            border_width=1
        )
        
        # 数字
        num_box = self.add_text_box(
            slide, left, top + Inches(0.2),
            width, Inches(0.6),
            str(number), font_size=36, font_color=self.colors['primary'],
            font_bold=True, align=PP_ALIGN.CENTER
        )
        
        # 标签
        label_box = self.add_text_box(
            slide, left, top + Inches(0.8),
            width, Inches(0.4),
            label, font_size=12, font_color=RgbColor(127, 140, 141),
            align=PP_ALIGN.CENTER
        )
        
        return card
    
    def add_job_card(self, slide, left, top, width, job_data):
        """添加岗位卡片"""
        card_height = Inches(2.2)
        
        # 卡片背景
        card = self.add_rectangle(
            slide, left, top, width, card_height,
            fill_color=self.colors['white'],
            border_color=RgbColor(225, 232, 237),
            border_width=1
        )
        
        # 匹配分数标签
        score = job_data.get('score', '')
        score_width = Inches(0.8)
        score_box = self.add_rectangle(
            slide, left + width - score_width - Inches(0.15), top + Inches(0.15),
            score_width, Inches(0.35),
            fill_color=self.colors['success']
        )
        
        score_text = self.add_text_box(
            slide, left + width - score_width - Inches(0.15), top + Inches(0.15),
            score_width, Inches(0.35),
            f"{score}分", font_size=12, font_color=self.colors['white'],
            font_bold=True, align=PP_ALIGN.CENTER
        )
        
        # 标题
        title_box = self.add_text_box(
            slide, left + Inches(0.15), top + Inches(0.15),
            width - score_width - Inches(0.4), Inches(0.4),
            job_data.get('title', ''), font_size=14, font_color=self.colors['secondary'],
            font_bold=True, align=PP_ALIGN.LEFT
        )
        
        # 详细信息
        details = job_data.get('details', [])
        detail_text = ' | '.join([f"{d[0]}: {d[1]}" for d in details[:4]])
        detail_box = self.add_text_box(
            slide, left + Inches(0.15), top + Inches(0.55),
            width - Inches(0.3), Inches(0.3),
            detail_text, font_size=10, font_color=self.colors['black'],
            align=PP_ALIGN.LEFT
        )
        
        # 技能标签
        skills = job_data.get('skills', [])
        skill_text = '技能: ' + ', '.join(skills) if skills else ''
        skill_box = self.add_text_box(
            slide, left + Inches(0.15), top + Inches(0.85),
            width - Inches(0.3), Inches(0.3),
            skill_text, font_size=10, font_color=self.colors['primary'],
            align=PP_ALIGN.LEFT
        )
        
        # 推荐理由
        reason = job_data.get('reason', '')
        if reason:
            reason_bg = self.add_rectangle(
                slide, left + Inches(0.15), top + Inches(1.15),
                width - Inches(0.3), Inches(0.5),
                fill_color=RgbColor(212, 237, 218)
            )
            # 左边框
            left_border = self.add_rectangle(
                slide, left + Inches(0.15), top + Inches(1.15),
                Pt(4), Inches(0.5),
                fill_color=self.colors['success']
            )
            
            reason_box = self.add_text_box(
                slide, left + Inches(0.25), top + Inches(1.2),
                width - Inches(0.5), Inches(0.45),
                reason, font_size=9, font_color=self.colors['black'],
                align=PP_ALIGN.LEFT
            )
        
        # 投递建议
        suggestion = job_data.get('suggestion', '')
        if suggestion:
            suggestion_box = self.add_text_box(
                slide, left + Inches(0.15), top + Inches(1.7),
                width - Inches(0.3), Inches(0.4),
                suggestion, font_size=9, font_color=RgbColor(85, 85, 85),
                align=PP_ALIGN.LEFT
            )
        
        return card
    
    def add_summary_box(self, slide, left, top, width, height, title, items):
        """添加摘要框（渐变背景）"""
        # 渐变背景
        bg = self.add_shape_with_gradient(
            slide, left, top, width, height,
            self.colors['gradient_start'], self.colors['gradient_end']
        )
        
        # 标题
        title_box = self.add_text_box(
            slide, left + Inches(0.2), top + Inches(0.15),
            width - Inches(0.4), Inches(0.4),
            title, font_size=18, font_color=self.colors['white'],
            font_bold=True, align=PP_ALIGN.LEFT
        )
        
        # 内容项
        y_offset = Inches(0.55)
        for item in items:
            item_box = self.add_text_box(
                slide, left + Inches(0.3), top + y_offset,
                width - Inches(0.5), Inches(0.35),
                f"• {item}", font_size=12, font_color=self.colors['white'],
                align=PP_ALIGN.LEFT
            )
            y_offset += Inches(0.35)
        
        return bg
    
    def add_info_box(self, slide, left, top, width, height, title, content, box_type='info'):
        """添加信息框"""
        # 背景颜色
        bg_colors = {
            'info': RgbColor(209, 236, 241),
            'warning': RgbColor(255, 243, 205),
            'success': RgbColor(212, 237, 218),
        }
        border_colors = {
            'info': self.colors['info'],
            'warning': self.colors['warning'],
            'success': self.colors['success'],
        }
        
        bg_color = bg_colors.get(box_type, bg_colors['info'])
        border_color = border_colors.get(box_type, border_colors['info'])
        
        # 背景
        bg = self.add_rectangle(
            slide, left, top, width, height,
            fill_color=bg_color
        )
        
        # 左边框
        left_border = self.add_rectangle(
            slide, left, top, Pt(4), height,
            fill_color=border_color
        )
        
        # 标题
        if title:
            title_box = self.add_text_box(
                slide, left + Inches(0.15), top + Inches(0.1),
                width - Inches(0.3), Inches(0.3),
                title, font_size=14, font_color=self.colors['black'],
                font_bold=True, align=PP_ALIGN.LEFT
            )
            content_top = top + Inches(0.4)
        else:
            content_top = top + Inches(0.1)
        
        # 内容
        if isinstance(content, list):
            y_offset = Inches(0)
            for item in content:
                item_box = self.add_text_box(
                    slide, left + Inches(0.2), content_top + y_offset,
                    width - Inches(0.4), Inches(0.25),
                    f"• {item}", font_size=11, font_color=self.colors['black'],
                    align=PP_ALIGN.LEFT
                )
                y_offset += Inches(0.25)
        else:
            content_box = self.add_text_box(
                slide, left + Inches(0.15), content_top,
                width - Inches(0.3), height - Inches(0.5),
                content, font_size=11, font_color=self.colors['black'],
                align=PP_ALIGN.LEFT
            )
        
        return bg
    
    def add_timeline_item(self, slide, left, top, width, label, content_title, items, box_type='info'):
        """添加时间线项"""
        item_height = Inches(0.25) * len(items) + Inches(0.6)
        
        # 时间标签
        label_box = self.add_text_box(
            slide, left, top,
            Inches(1.2), Inches(0.3),
            label, font_size=12, font_color=self.colors['primary'],
            font_bold=True, align=PP_ALIGN.LEFT
        )
        
        # 内容框
        content_left = left + Inches(1.3)
        content_width = width - Inches(1.3)
        
        self.add_info_box(
            slide, content_left, top,
            content_width, item_height,
            content_title, items, box_type
        )
        
        return item_height
    
    def parse_table(self, table_elem):
        """解析HTML表格"""
        table_data = []
        
        # 解析表头
        thead = table_elem.find('thead')
        if thead:
            for tr in thead.find_all('tr'):
                row_data = []
                for th in tr.find_all('th'):
                    row_data.append({'text': th.get_text(strip=True)})
                if row_data:
                    table_data.append(row_data)
        
        # 解析表体
        tbody = table_elem.find('tbody')
        if tbody:
            for tr in tbody.find_all('tr'):
                row_data = []
                for td in tr.find_all('td'):
                    cell_text = td.get_text(strip=True)
                    cell_data = {'text': cell_text}
                    
                    # 检查是否有badge
                    badge = td.find(class_='badge')
                    if badge:
                        cell_data['badge'] = True
                        if 'badge-high' in badge.get('class', []):
                            cell_data['badge_type'] = 'high'
                        elif 'badge-low' in badge.get('class', []):
                            cell_data['badge_type'] = 'low'
                        else:
                            cell_data['badge_type'] = 'medium'
                    
                    row_data.append(cell_data)
                if row_data:
                    table_data.append(row_data)
        
        # 如果没有thead/tbody，直接解析tr
        if not table_data:
            for tr in table_elem.find_all('tr'):
                row_data = []
                for cell in tr.find_all(['th', 'td']):
                    row_data.append({'text': cell.get_text(strip=True)})
                if row_data:
                    table_data.append(row_data)
        
        return table_data
    
    def parse_job_card(self, card_elem):
        """解析岗位卡片"""
        job_data = {}
        
        # 匹配分数
        score_elem = card_elem.find(class_='match-score')
        if score_elem:
            score_text = score_elem.get_text(strip=True)
            score_match = re.search(r'(\d+)', score_text)
            if score_match:
                job_data['score'] = score_match.group(1)
        
        # 标题
        title_elem = card_elem.find('h4')
        if title_elem:
            job_data['title'] = title_elem.get_text(strip=True)
        
        # 详细信息表格
        table = card_elem.find('table')
        if table:
            details = []
            for tr in table.find_all('tr'):
                cells = tr.find_all('td')
                if len(cells) >= 2:
                    details.append((cells[0].get_text(strip=True), cells[1].get_text(strip=True)))
            job_data['details'] = details
        
        # 技能标签
        skills = []
        for skill in card_elem.find_all(class_='skill-tag'):
            skills.append(skill.get_text(strip=True))
        job_data['skills'] = skills
        
        # 推荐理由
        recommendation = card_elem.find(class_='recommendation')
        if recommendation:
            job_data['reason'] = recommendation.get_text(strip=True)
        
        # 投递建议
        suggestion_p = card_elem.find('p')
        if suggestion_p and '投递建议' in suggestion_p.get_text():
            job_data['suggestion'] = suggestion_p.get_text(strip=True)
        
        return job_data
    
    def convert(self):
        """执行转换"""
        container = self.soup.find(class_='container')
        if not container:
            container = self.soup.find('body')
        
        # 1. 创建标题页
        title_elem = self.soup.find('h1')
        title_text = title_elem.get_text(strip=True) if title_elem else 'HTML Report'
        
        meta_elem = self.soup.find(class_='meta')
        meta_text = ''
        if meta_elem:
            meta_lines = [p.get_text(strip=True) for p in meta_elem.find_all('p')]
            meta_text = '\n'.join(meta_lines[:3])
        
        self.create_title_slide(title_text, meta_text)
        
        # 2. 核心结论页
        summary_box = self.soup.find(class_='summary-box')
        if summary_box:
            slide = self.create_content_slide('核心结论', [])
            
            items = []
            for li in summary_box.find_all('li'):
                items.append(li.get_text(strip=True))
            
            self.add_summary_box(
                slide, self.margin_left, Inches(1.2),
                self.content_width, Inches(2.5),
                '核心结论', items
            )
        
        # 3. 统计数据概览页
        stat_cards = self.soup.find_all(class_='stat-card')
        if stat_cards and len(stat_cards) >= 4:
            slide = self.create_content_slide('一、统计数据概览', [])
            
            card_width = Inches(2.8)
            card_height = Inches(1.5)
            gap = Inches(0.3)
            start_left = self.margin_left + Inches(0.2)
            
            for i, card in enumerate(stat_cards[:4]):
                h4 = card.find('h4')
                p = card.find('p')
                number = h4.get_text(strip=True) if h4 else ''
                label = p.get_text(strip=True) if p else ''
                
                left = start_left + (card_width + gap) * i
                self.add_stat_card(slide, left, Inches(2), card_width, card_height, number, label)
        
        # 4. 岗位检索概况页 - 经验分布图
        chart_containers = self.soup.find_all(class_='chart-container')
        if chart_containers:
            for chart in chart_containers:
                img = chart.find('img')
                h3 = chart.find('h3')
                
                if img and img.get('src'):
                    chart_title = h3.get_text(strip=True) if h3 else '图表'
                    slide = self.create_content_slide(f'二、岗位检索概况 - {chart_title}', [])
                    
                    img_path = img.get('src')
                    self.add_image_to_slide(
                        slide, img_path,
                        Inches(2), Inches(1.3),
                        width=Inches(9)
                    )
        
        # 5. 经验要求分布表格页
        tables = self.soup.find_all('table')
        table_idx = 0
        for table in tables:
            # 跳过job-card内的表格
            if table.find_parent(class_='job-card'):
                continue
            
            table_data = self.parse_table(table)
            if table_data and len(table_data) > 1:
                # 查找表格前的标题
                prev_h3 = table.find_previous_sibling('h3')
                prev_h2 = table.find_previous('h2')
                
                table_title = ''
                if prev_h3:
                    table_title = prev_h3.get_text(strip=True)
                elif prev_h2:
                    table_title = prev_h2.get_text(strip=True)
                
                if not table_title:
                    table_title = f'数据表格 {table_idx + 1}'
                
                slide = self.create_content_slide(table_title, [])
                
                # 计算表格宽度
                num_cols = len(table_data[0])
                table_width = min(self.content_width, Inches(num_cols * 2.5))
                
                self.add_table_to_slide(
                    slide, table_data,
                    self.margin_left, Inches(1.3),
                    table_width
                )
                
                table_idx += 1
        
        # 6. 高匹配度岗位推荐页
        job_cards = self.soup.find_all(class_='job-card')
        job_cards_data = []
        
        for card in job_cards:
            # 检查是否是真正的岗位卡片（有match-score）
            if card.find(class_='match-score'):
                job_data = self.parse_job_card(card)
                if job_data.get('title'):
                    job_cards_data.append(job_data)
        
        # 每页显示2个岗位卡片
        if job_cards_data:
            for i in range(0, len(job_cards_data), 2):
                slide = self.create_content_slide(
                    f'三、高匹配度岗位推荐（{i+1}-{min(i+2, len(job_cards_data))}）', []
                )
                
                card_width = Inches(6)
                
                for j, job_data in enumerate(job_cards_data[i:i+2]):
                    left = self.margin_left + (card_width + Inches(0.3)) * j
                    self.add_job_card(slide, left, Inches(1.3), card_width, job_data)
        
        # 7. 技能要求分析页
        skill_section = None
        for h2 in self.soup.find_all('h2'):
            if '技能要求' in h2.get_text():
                skill_section = h2
                break
        
        if skill_section:
            skill_table = skill_section.find_next('table')
            if skill_table:
                table_data = self.parse_table(skill_table)
                if table_data:
                    slide = self.create_content_slide('四、技能要求分析', [])
                    self.add_table_to_slide(
                        slide, table_data,
                        self.margin_left, Inches(1.3),
                        self.content_width
                    )
        
        # 8. 技能缺口与补齐计划页
        timeline = self.soup.find(class_='timeline')
        if timeline:
            slide = self.create_content_slide('五、技能缺口与补齐计划', [])
            
            timeline_items = timeline.find_all(class_='timeline-item')
            y_pos = Inches(1.3)
            
            for item in timeline_items:
                label_elem = item.find(class_='timeline-label')
                content_elem = item.find(class_='timeline-content')
                
                if label_elem and content_elem:
                    label = label_elem.get_text(strip=True)
                    
                    # 获取内容类型和标题
                    box_type = 'info'
                    if content_elem.find(class_='warning'):
                        box_type = 'warning'
                    elif content_elem.find(class_='recommendation'):
                        box_type = 'success'
                    
                    strong = content_elem.find('strong')
                    content_title = strong.get_text(strip=True) if strong else ''
                    
                    items = []
                    for li in content_elem.find_all('li'):
                        items.append(li.get_text(strip=True))
                    
                    if items:
                        item_height = self.add_timeline_item(
                            slide, self.margin_left, y_pos,
                            self.content_width, label, content_title, items, box_type
                        )
                        y_pos += item_height + Inches(0.15)
        
        # 9. 求职策略建议页
        strategy_section = None
        for h2 in self.soup.find_all('h2'):
            if '求职策略' in h2.get_text():
                strategy_section = h2
                break
        
        if strategy_section:
            slide = self.create_content_slide('六、求职策略建议', [])
            
            # 找到优先级卡片
            section_grid = strategy_section.find_next(class_='section-grid')
            if section_grid:
                priority_cards = section_grid.find_all(class_='job-card')
                
                card_width = Inches(4)
                gap = Inches(0.15)
                
                for i, card in enumerate(priority_cards[:3]):
                    h4 = card.find('h4')
                    title = h4.get_text(strip=True) if h4 else f'优先级{i+1}'
                    
                    items = []
                    for li in card.find_all('li'):
                        items.append(li.get_text(strip=True))
                    
                    left = self.margin_left + (card_width + gap) * i
                    
                    # 卡片背景
                    card_bg = self.add_rectangle(
                        slide, left, Inches(1.3),
                        card_width, Inches(2.5),
                        fill_color=self.colors['white'],
                        border_color=RgbColor(225, 232, 237),
                        border_width=1
                    )
                    
                    # 标题
                    title_box = self.add_text_box(
                        slide, left + Inches(0.1), Inches(1.4),
                        card_width - Inches(0.2), Inches(0.4),
                        title, font_size=14, font_color=self.colors['secondary'],
                        font_bold=True, align=PP_ALIGN.LEFT
                    )
                    
                    # 列表项
                    y_offset = Inches(1.8)
                    for item in items:
                        item_box = self.add_text_box(
                            slide, left + Inches(0.15), y_offset,
                            card_width - Inches(0.3), Inches(0.35),
                            item, font_size=10, font_color=self.colors['black'],
                            align=PP_ALIGN.LEFT
                        )
                        y_offset += Inches(0.35)
            
            # 时间安排建议
            info_box = strategy_section.find_next(class_='info')
            if info_box:
                h3 = info_box.find('h3')
                title = h3.get_text(strip=True) if h3 else ''
                
                items = []
                for li in info_box.find_all('li'):
                    items.append(li.get_text(strip=True))
                
                self.add_info_box(
                    slide, self.margin_left, Inches(4),
                    self.content_width, Inches(1.5),
                    title, items, 'info'
                )
        
        # 10. 结论与展望页
        conclusion_section = None
        for h2 in self.soup.find_all('h2'):
            if '结论' in h2.get_text():
                conclusion_section = h2
                break
        
        if conclusion_section:
            slide = self.create_content_slide('七、结论与展望', [])
            
            # 核心结论
            summary_box = conclusion_section.find_next(class_='summary-box')
            if summary_box:
                items = []
                for li in summary_box.find_all('li'):
                    items.append(li.get_text(strip=True))
                
                self.add_summary_box(
                    slide, self.margin_left, Inches(1.2),
                    self.content_width, Inches(2.2),
                    '核心结论', items
                )
            
            # 职业发展路径
            section_grid = conclusion_section.find_next(class_='section-grid')
            if section_grid:
                stat_cards = section_grid.find_all(class_='stat-card')
                
                card_width = Inches(2.8)
                card_height = Inches(1.3)
                gap = Inches(0.3)
                
                for i, card in enumerate(stat_cards[:4]):
                    h4 = card.find('h4')
                    p = card.find('p')
                    number = h4.get_text(strip=True) if h4 else ''
                    label = p.get_text(separator='\n', strip=True) if p else ''
                    
                    left = self.margin_left + Inches(0.2) + (card_width + gap) * i
                    
                    # 卡片背景
                    card_bg = self.add_rectangle(
                        slide, left, Inches(3.6),
                        card_width, card_height,
                        fill_color=self.colors['white'],
                        border_color=RgbColor(225, 232, 237),
                        border_width=1
                    )
                    
                    # 数字
                    num_box = self.add_text_box(
                        slide, left, Inches(3.7),
                        card_width, Inches(0.5),
                        number, font_size=24, font_color=self.colors['primary'],
                        font_bold=True, align=PP_ALIGN.CENTER
                    )
                    
                    # 标签
                    label_box = self.add_text_box(
                        slide, left, Inches(4.2),
                        card_width, Inches(0.6),
                        label, font_size=10, font_color=RgbColor(127, 140, 141),
                        align=PP_ALIGN.CENTER
                    )
        
        # 11. 报告说明页
        report_info = None
        for info_box in self.soup.find_all(class_='info'):
            h3 = info_box.find('h3')
            if h3 and '报告说明' in h3.get_text():
                report_info = info_box
                break
        
        if report_info:
            slide = self.create_content_slide('报告说明', [])
            
            items = []
            for p in report_info.find_all('p'):
                items.append(p.get_text(strip=True))
            
            self.add_info_box(
                slide, self.margin_left, Inches(1.3),
                self.content_width, Inches(2),
                '', items, 'info'
            )
        
        # 12. 感谢页
        slide = self.add_blank_slide()
        
        # 渐变背景
        bg = self.add_shape_with_gradient(
            slide, Inches(0), Inches(0),
            self.slide_width, self.slide_height,
            self.colors['gradient_start'], self.colors['gradient_end']
        )
        
        # 感谢文字
        thanks_box = self.add_text_box(
            slide, self.margin_left, Inches(3),
            self.content_width, Inches(1.5),
            '感谢阅读', font_size=48, font_color=self.colors['white'],
            font_bold=True, align=PP_ALIGN.CENTER
        )
        
        # 保存文件
        self.prs.save(self.output_path)
        print(f"PPTX saved to: {self.output_path}")
        
        return self.output_path


def main():
    """主函数"""
    if len(sys.argv) < 2:
        print("Usage: python html_to_pptx.py <input.html> [output.pptx]")
        sys.exit(1)
    
    html_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else None
    
    if not os.path.exists(html_path):
        print(f"Error: File not found: {html_path}")
        sys.exit(1)
    
    converter = HTMLToPPTXConverter(html_path, output_path)
    output = converter.convert()
    print(f"Conversion completed: {output}")


if __name__ == '__main__':
    main()
