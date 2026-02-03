"""Sandbox-only tools for media processing and conversions."""

from __future__ import annotations

import base64
import json
import textwrap
import uuid
from pathlib import Path
from typing import Any

from langchain_core.tools import BaseTool, tool

from deepagents.backends.protocol import SandboxBackendProtocol

from .sandbox_tool_utils import (
    _ensure_workspace_path,
    _run_python_in_sandbox,
    _sync_to_host,
)


def build_sandbox_tools(
    *, sandbox_backend: SandboxBackendProtocol, workspace_dir: Path | None = None
) -> list[BaseTool]:
    """Build sandbox-backed tools that execute inside the container."""

    @tool(
        "pdf_to_word",
        description=(
            "Convert a PDF to DOCX inside the sandbox. Provide pdf_path and optional output_path."
        ),
    )
    def pdf_to_word(
        pdf_path: str,
        output_path: str | None = None,
    ) -> dict[str, Any]:
        pdf_path = _ensure_workspace_path(pdf_path)
        if not output_path:
            stem = Path(pdf_path).stem or f"document_{uuid.uuid4().hex[:6]}"
            output_path = f"/workspace/converted/{stem}.docx"
        output_path = _ensure_workspace_path(output_path)

        args = {"pdf_path": pdf_path, "output_path": output_path}
        args_b64 = base64.b64encode(json.dumps(args).encode("utf-8")).decode("ascii")

        script = textwrap.dedent(
            """
            import base64
            import json
            import os
            from pdf2docx import Converter

            args = json.loads(base64.b64decode("__ARGS_B64__").decode("utf-8"))
            pdf_path = args.get("pdf_path", "")
            output_path = args.get("output_path", "")

            if not os.path.isfile(pdf_path):
                print("RESULT_JSON:" + json.dumps({"success": False, "error": "PDF not found: " + str(pdf_path)}, ensure_ascii=False))
                raise SystemExit(0)

            os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
            try:
                converter = Converter(pdf_path)
                converter.convert(output_path)
                converter.close()
            except Exception as exc:
                print("RESULT_JSON:" + json.dumps({"success": False, "error": str(exc)}, ensure_ascii=False))
                raise SystemExit(0)

            print("RESULT_JSON:" + json.dumps({"success": True, "output_path": output_path}, ensure_ascii=False))
            """
        )
        script = script.replace("__ARGS_B64__", args_b64)
        result = _run_python_in_sandbox(sandbox_backend, script)
        output_path = result.get("output_path")
        if isinstance(output_path, str):
            host_path = _sync_to_host(sandbox_backend, workspace_dir, output_path)
            if host_path:
                result["host_path"] = host_path
        return result

    @tool(
        "extract_file_text",
        description=(
            "Extract text from a file inside the sandbox. "
            "Supports txt/md/csv/json/yaml, pdf, docx, xlsx/xls. "
            "Provide file_path (inside /workspace) and optional max_chars."
        ),
    )
    def extract_file_text(
        file_path: str,
        max_chars: int = 50000,
    ) -> dict[str, Any]:
        file_path = _ensure_workspace_path(file_path)
        args = {"file_path": file_path, "max_chars": max_chars}
        args_b64 = base64.b64encode(json.dumps(args).encode("utf-8")).decode("ascii")

        script = textwrap.dedent(
            """
            import base64
            import json
            import os
            import zipfile
            import xml.etree.ElementTree as ET
            from pathlib import Path

            args = json.loads(base64.b64decode("__ARGS_B64__").decode("utf-8"))
            file_path = args.get("file_path", "")
            max_chars = int(args.get("max_chars") or 50000)

            def _truncate(text):
                if len(text) <= max_chars:
                    return text, False
                return text[:max_chars] + "\\n... (truncated)", True

            def _read_text(path):
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()

            def _extract_docx(path):
                with zipfile.ZipFile(path) as z:
                    if "word/document.xml" not in z.namelist():
                        return ""
                    xml_bytes = z.read("word/document.xml")
                root = ET.fromstring(xml_bytes)
                parts = []
                for elem in root.iter():
                    tag = elem.tag
                    if tag.endswith("}t"):
                        if elem.text:
                            parts.append(elem.text)
                    elif tag.endswith("}tab"):
                        parts.append("\\t")
                    elif tag.endswith("}br") or tag.endswith("}p"):
                        parts.append("\\n")
                return "".join(parts)

            def _extract_pdf(path):
                from pypdf import PdfReader

                reader = PdfReader(path)
                pages = []
                for page in reader.pages:
                    text = page.extract_text() or ""
                    if text:
                        pages.append(text)
                return "\\n\\n".join(pages)

            def _extract_xlsx(path):
                from openpyxl import load_workbook

                wb = load_workbook(path, data_only=True)
                chunks = []
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    rows = []
                    for row in ws.iter_rows(values_only=True):
                        rows.append("\\t".join("" if v is None else str(v) for v in row))
                    sheet_text = "\\n".join(rows).strip()
                    chunks.append("## Sheet: " + sheet_name + "\\n" + sheet_text)
                return "\\n\\n".join(chunks)

            def _extract_xls(path):
                try:
                    import xlrd
                except Exception as exc:
                    raise RuntimeError("xlrd not installed: " + str(exc))

                book = xlrd.open_workbook(path)
                chunks = []
                for sheet in book.sheets():
                    rows = []
                    for r in range(sheet.nrows):
                        rows.append("\\t".join(str(sheet.cell_value(r, c)) for c in range(sheet.ncols)))
                    sheet_text = "\\n".join(rows).strip()
                    chunks.append("## Sheet: " + sheet.name + "\\n" + sheet_text)
                return "\\n\\n".join(chunks)

            if not os.path.isfile(file_path):
                print("RESULT_JSON:" + json.dumps({"success": False, "error": "File not found: " + str(file_path), "file_path": file_path}, ensure_ascii=False))
                raise SystemExit(0)

            ext = Path(file_path).suffix.lower()
            text = ""
            try:
                if ext in {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".log", ".ini", ".cfg"}:
                    text = _read_text(file_path)
                elif ext == ".pdf":
                    text = _extract_pdf(file_path)
                elif ext == ".docx":
                    text = _extract_docx(file_path)
                elif ext == ".xlsx":
                    text = _extract_xlsx(file_path)
                elif ext == ".xls":
                    text = _extract_xls(file_path)
                else:
                    raise RuntimeError("Unsupported file type: " + (ext or "unknown"))
            except Exception as exc:
                print("RESULT_JSON:" + json.dumps({"success": False, "error": str(exc), "file_path": file_path}, ensure_ascii=False))
                raise SystemExit(0)

            text, truncated = _truncate(text)
            print("RESULT_JSON:" + json.dumps({"success": True, "text": text, "truncated": truncated, "file_path": file_path}, ensure_ascii=False))
            """
        )
        script = script.replace("__ARGS_B64__", args_b64)
        return _run_python_in_sandbox(sandbox_backend, script)

    return [pdf_to_word, extract_file_text]


__all__ = ["build_sandbox_tools"]
