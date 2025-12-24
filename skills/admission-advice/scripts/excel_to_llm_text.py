
import argparse
from openpyxl import load_workbook


def fill_merged_cells(ws):
    """
    将合并单元格拆开，并把左上角的值填充到每一个单元格
    """
    merged_ranges = list(ws.merged_cells.ranges)

    for merged_range in merged_ranges:
        min_row, min_col, max_row, max_col = (
            merged_range.min_row,
            merged_range.min_col,
            merged_range.max_row,
            merged_range.max_col,
        )

        value = ws.cell(row=min_row, column=min_col).value

        ws.unmerge_cells(str(merged_range))

        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(row=r, column=c).value = value


def excel_to_text(input_path, output_path):
    wb = load_workbook(input_path, data_only=True)

    with open(output_path, "w", encoding="utf-8") as f:
        for ws in wb.worksheets:
            fill_merged_cells(ws)

            for row in ws.iter_rows():
                values = []
                for cell in row:
                    if cell.value is None:
                        values.append("")
                    else:
                        values.append(str(cell.value))
                f.write("|".join(values) + "\n")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("-i", "--input", required=True, help="input excel file")
    parser.add_argument("-o", "--output", required=True, help="output text file")
    args = parser.parse_args()

    excel_to_text(args.input, args.output)


if __name__ == "__main__":
    main()
